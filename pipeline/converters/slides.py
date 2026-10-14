"""PowerPoint decks (`.pptx`). Google Slides arrive as a `.pptx` exported from Slides.

Legacy `.ppt` and OpenDocument `.odp` decks are turned into a `.pptx` by LibreOffice first,
as a subprocess into `WORK_DIR`, the same route legacy Word files take.

Each slide is one page: its blocks are anchored `pNNN` with NNN the slide number, and the
slide opens with `## Slide N: <title>` so a reader (and the sectioniser) can find it by the
number people use for it.

The slides are walked here with python-pptx rather than through Docling's own PowerPoint
walk. Measured on the fixture deck (docling-slim 2.126): Docling put every bullet at one
level, so a sub-point read as a point of its own, and it carries no hidden-slide flag and no
SmartArt text. Two of its helpers are still used, because they are the hard part and are
tested upstream: `_is_list_item` (whether a paragraph is a bullet, resolved through the
slide layout and master) and `_iter_shapes_by_position` (reading order by position on the
slide rather than by creation order). Both are private, so a Docling upgrade that changes
them fails `tests/test_slides.py` rather than the output quietly changing.

What a deck carries beyond its slide text, and what happens to it:

* **Speaker notes** are kept, as a quoted block at the end of their slide.
* **Reviewer comments** are left out, as they are for Word: they are review traffic, not
  the deck's content.
* **Hidden slides** are kept and marked.
* **Charts** keep their plotted numbers as a table; the numbers are in the file.
* **SmartArt** keeps its text in a separate part; it is read from there as a list, placed
  after the slide's other shapes because its position cannot be ordered with them.
* **Pictures** are marked in place and counted in an INCOMPLETE note, as for Word.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from ..config import Config
from . import render_paged_block_provenance
from .csv_table import escape_cell
from .office import PICTURE_MARKER, _sanitise_package, soffice_convert

CONVERTER_PPTX = "python-pptx (slides)"
CONVERTER_LIBREOFFICE_PPTX = "libreoffice+python-pptx (slides)"

#: Deck formats LibreOffice turns into `.pptx` before conversion.
LEGACY_SUFFIXES = (".ppt", ".odp")

HIDDEN_MARKER = "> **Hidden slide:** not shown when the deck is presented."

_DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "dgm": _DIAGRAM_URI,
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
_R_DM = "{%s}dm" % _NS["r"]

_logger = logging.getLogger(__name__)


def _load(path: Path):
    """Open the deck through Docling's backend, for its helpers. Its `convert()` is not run."""
    from docling.backend.mspowerpoint_backend import (  # noqa: PLC0415 - lazy by design
        MsPowerpointDocumentBackend,
    )
    from docling.datamodel.base_models import InputFormat  # noqa: PLC0415
    from docling.datamodel.document import InputDocument  # noqa: PLC0415

    in_doc = InputDocument(
        path_or_stream=path,
        format=InputFormat.PPTX,
        backend=MsPowerpointDocumentBackend,
        filename=path.name,
    )
    sanitised, dangling = _sanitise_package(path)
    if dangling:
        _logger.warning(
            "%s: %d relationship(s) point at parts the package does not contain; "
            "treated as external references",
            path.name,
            len(dangling),
        )
    backend = MsPowerpointDocumentBackend(
        in_doc=in_doc, path_or_stream=sanitised if sanitised is not None else path
    )
    return backend, dangling


def _one_line(text: str) -> str:
    return " ".join(text.replace("\v", " ").split())


def _cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _table_markdown(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(escape_cell(cell) for cell in rows[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")
    lines.extend("| " + " | ".join(escape_cell(cell) for cell in row) + " |" for row in rows[1:])
    return "\n".join(lines)


def _table_block(shape) -> tuple[str, str, str] | None:
    rows = [
        ["" if cell.is_spanned else _one_line(cell.text) for cell in row.cells]
        for row in shape.table.rows
    ]
    if not any(any(cell for cell in row) for row in rows):
        return None
    return _table_markdown(rows), "table", "structural"


def _chart_block(shape) -> tuple[str, str, str] | None:
    chart = shape.chart
    title = ""
    if chart.has_title and chart.chart_title.has_text_frame:
        title = _one_line(chart.chart_title.text_frame.text)
    label = f"**Chart:** {title}" if title else "**Chart**"
    plots = list(chart.plots)
    categories = [_cell_value(c) for c in plots[0].categories] if plots else []
    series = list(chart.series)
    if not series:
        return f"{label} (the file holds no data for it)", "paragraph", "pptx:chart"
    columns = [(_cell_value(s.name), [_cell_value(v) for v in s.values]) for s in series]
    height = max([len(categories)] + [len(values) for _, values in columns])
    rows = [[""] + [name for name, _ in columns]]
    for index in range(height):
        row = [categories[index] if index < len(categories) else ""]
        row += [values[index] if index < len(values) else "" for _, values in columns]
        rows.append(row)
    return f"{label}\n\n{_table_markdown(rows)}", "table", "pptx:chart"


def _smartart_items(shape) -> list[str] | None:
    """The text of each SmartArt node, in the order the data part lists them.

    `None` when the frame is not SmartArt; an empty list when it is but its data part is
    missing or unreadable.
    """
    graphic_data = shape._element.find("a:graphic/a:graphicData", _NS)
    if graphic_data is None or graphic_data.get("uri") != _DIAGRAM_URI:
        return None
    rel_ids = graphic_data.find("dgm:relIds", _NS)
    r_id = rel_ids.get(_R_DM) if rel_ids is not None else None
    try:
        from lxml import etree  # noqa: PLC0415

        part = shape.part.related_part(r_id)
        root = etree.fromstring(part.blob)
    except Exception:  # noqa: BLE001 - any failure here means the text is not recoverable
        return []
    items = []
    for point in root.iterfind("dgm:ptLst/dgm:pt", _NS):
        if point.get("type", "node") != "node":
            continue
        for paragraph in point.iterfind("dgm:t/a:p", _NS):
            text = _one_line("".join(t.text or "" for t in paragraph.iterfind(".//a:t", _NS)))
            if text:
                items.append(text)
    return items


class _SlideWalk:
    """Blocks of one slide, in reading order, and the counts the deck-level notes need."""

    def __init__(self, backend, counts: dict[str, int]) -> None:
        self.backend = backend
        self.counts = counts
        self.blocks: list[tuple[str, str, str]] = []
        self.smartart: list[tuple[str, str, str]] = []

    def shapes(self, shapes, skip_id: int | None) -> None:
        from pptx.shapes.group import GroupShape  # noqa: PLC0415
        from pptx.shapes.picture import Picture  # noqa: PLC0415

        for shape in self.backend._iter_shapes_by_position(shapes):
            if skip_id is not None and shape.shape_id == skip_id:
                continue
            if isinstance(shape, GroupShape):
                self.shapes(shape.shapes, skip_id)
                continue
            if isinstance(shape, Picture):
                self.counts["images"] += 1
                self.blocks.append((PICTURE_MARKER, "picture", "structural"))
                continue
            if getattr(shape, "has_table", False):
                block = _table_block(shape)
                if block:
                    self.blocks.append(block)
                continue
            if getattr(shape, "has_chart", False):
                self.counts["charts"] += 1
                block = _chart_block(shape)
                if block:
                    self.blocks.append(block)
                continue
            if shape._element.tag.endswith("}graphicFrame"):
                items = _smartart_items(shape)
                if items is not None:
                    if items:
                        self.smartart.append(
                            ("\n".join(f"- {item}" for item in items), "list", "pptx:smartart")
                        )
                    else:
                        self.counts["unreadable_smartart"] += 1
                continue
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                self.text_frame(shape)

    def text_frame(self, shape) -> None:
        from pptx.enum.shapes import PP_PLACEHOLDER  # noqa: PLC0415

        is_subtitle = (
            shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE
        )
        list_lines: list[str] = []
        numbers: dict[int, int] = {}

        def flush_list() -> None:
            if list_lines:
                self.blocks.append(("\n".join(list_lines), "list", "structural"))
            list_lines.clear()
            numbers.clear()

        for paragraph in shape.text_frame.paragraphs:
            text = _one_line(paragraph.text)
            if not text:
                continue
            is_list, bullet_type = self.backend._is_list_item(paragraph)
            if is_list and not is_subtitle:
                level = paragraph.level
                for deeper in [lvl for lvl in numbers if lvl > level]:
                    del numbers[deeper]
                if bullet_type == "Numbered":
                    numbers[level] = numbers.get(level, 0) + 1
                    marker = f"{numbers[level]}."
                else:
                    marker = "-"
                list_lines.append(f"{'  ' * level}{marker} {text}")
                continue
            flush_list()
            if is_subtitle:
                # Not a heading: a heading with nothing under it is merged into the next
                # section, which put slide 2's content under slide 1's subtitle.
                self.blocks.append((text, "paragraph", "pptx:subtitle"))
            else:
                self.blocks.append((text, "paragraph", "structural"))
        flush_list()


def _notes_block(slide) -> tuple[str, str, str] | None:
    if not slide.has_notes_slide:
        return None
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return None
    lines = [line.rstrip() for line in frame.text.replace("\v", "\n").strip().split("\n")]
    if not any(lines):
        return None
    quoted = [f"> **Speaker notes:** {lines[0]}"]
    quoted += [f"> {line}" if line else ">" for line in lines[1:]]
    return "\n".join(quoted), "paragraph", "pptx:notes"


def convert_with_provenance(
    path: Path, config: Config, *, provenance_slug: str
) -> tuple[str, str, list[dict[str, Any]], dict[str, Any]]:
    """Convert a `.pptx`. Returns `(body, converter, provenance, extras)`.

    `extras` records `images`, `charts`, `hidden_slides`, `unreadable_smartart` and
    `dangling_relationships` for the manifest and the report.
    """
    converter = CONVERTER_PPTX
    if path.suffix.lower() in LEGACY_SUFFIXES:
        path = soffice_convert(path, config, "pptx", "slides2pptx")
        converter = CONVERTER_LIBREOFFICE_PPTX
    backend, dangling = _load(path)
    presentation = backend.pptx_obj
    counts = {"images": 0, "charts": 0, "hidden_slides": 0, "unreadable_smartart": 0}

    pages: list[tuple[int, list[tuple[str, str, str]]]] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = _one_line(title_shape.text) if title_shape is not None else ""
        heading = f"## Slide {number}: {title}" if title else f"## Slide {number}"
        blocks = [(heading, "heading", "pptx:slide")]
        if slide._element.get("show") in ("0", "false"):
            counts["hidden_slides"] += 1
            blocks.append((HIDDEN_MARKER, "paragraph", "pptx:hidden"))

        walk = _SlideWalk(backend, counts)
        walk.shapes(slide.shapes, title_shape.shape_id if title and title_shape else None)
        blocks += walk.blocks + walk.smartart
        notes = _notes_block(slide)
        if notes:
            blocks.append(notes)
        pages.append((number, blocks))

    body, provenance = render_paged_block_provenance(pages, provenance_slug)

    notes: list[str] = []
    if counts["images"]:
        notes.append(
            f"> **INCOMPLETE — this deck shows {counts['images']} image(s), and their "
            "content is not in the text.** No OCR was attempted. Each is marked in place "
            "below."
        )
    if counts["unreadable_smartart"]:
        notes.append(
            f"> **INCOMPLETE — {counts['unreadable_smartart']} SmartArt diagram(s) in this "
            "deck could not be read.** Their text is missing below."
        )
    if dangling:
        notes.append(
            f"> **INCOMPLETE — {len(dangling)} reference(s) in this deck point at content "
            "the file does not contain** (for example an image that was never packaged). "
            "Nothing can be recovered for them."
        )
    if notes:
        body = "\n\n".join([*notes, body])

    extras = {**counts, "dangling_relationships": len(dangling)}
    return body, converter, provenance, extras
